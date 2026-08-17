"""Generic Excel reading. These SWE.1/supporting-doc workbooks don't reliably
have a header on row 1 — sometimes there are title/metadata rows above the
real table, sometimes there's no discoverable header at all — so everything
here works on a raw grid first and figures out structure afterward, rather
than assuming pandas-style clean tables.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

MIN_HEADER_NONEMPTY_CELLS = 3
HEADER_SEARCH_WINDOW = 40
MAX_HEADER_CELL_LEN = 40


@dataclass
class RawSheet:
    """One worksheet as a plain grid. rows[0] is Excel row 1, each row's
    index 0 is column A. Merged cells are flattened (top-left value
    propagated across the whole range) so downstream code sees a consistent
    value everywhere in that range instead of None.
    """

    name: str
    source_file: str
    rows: list[list[Any]] = field(default_factory=list)

    def iter_rows(self):
        """Yield (row_index, values), 1-indexed to match Excel."""
        for i, row in enumerate(self.rows, start=1):
            yield i, row


def load_workbook_sheets(path: Path) -> list[RawSheet]:
    """Load every sheet of an .xlsx/.xlsm file into RawSheets."""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return [_load_sheet(ws, source_file=path.name) for ws in wb.worksheets]
    finally:
        wb.close()


def _load_sheet(ws, source_file: str) -> RawSheet:
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    _flatten_merged_cells(ws, rows)
    return RawSheet(name=ws.title, source_file=source_file, rows=rows)


def _flatten_merged_cells(ws, rows: list[list[Any]]) -> None:
    for merged_range in ws.merged_cells.ranges:
        top_left_value = _safe_get(rows, merged_range.min_row, merged_range.min_col)
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                _safe_set(rows, r, c, top_left_value)


def _safe_get(rows: list[list[Any]], row: int, col: int) -> Any:
    if row - 1 >= len(rows) or col - 1 >= len(rows[row - 1]):
        return None
    return rows[row - 1][col - 1]


def _safe_set(rows: list[list[Any]], row: int, col: int, value: Any) -> None:
    if row - 1 >= len(rows) or col - 1 >= len(rows[row - 1]):
        return
    rows[row - 1][col - 1] = value


def _row_looks_like_header(row_values: list[Any]) -> bool:
    """A header row is (almost) always pure text - a data row with a numeric
    id/value mixed in among a few short text cells can otherwise look just
    like a header by cell count alone, so any non-string, non-empty cell
    disqualifies the row outright.
    """
    nonempty = [v for v in row_values if v not in (None, "")]
    if not nonempty or any(not isinstance(v, str) for v in nonempty):
        return False
    short_strings = [v for v in nonempty if v.strip() and len(v.strip()) < MAX_HEADER_CELL_LEN]
    return len(short_strings) >= MIN_HEADER_NONEMPTY_CELLS


def find_header_row(
    sheet: RawSheet, anchor_row: int | None = None
) -> tuple[int | None, list[str] | None]:
    """Find the header row for a sheet whose header isn't reliably on row 1.

    With an `anchor_row` (a known data row, e.g. a matched requirement row),
    scans upward from just above it — this is how we find a header that sits
    somewhere above a block of matched rows. Without one, scans downward from
    the top of the sheet — used for supporting-doc sheets with no such
    anchor. Both directions give up after HEADER_SEARCH_WINDOW rows and
    return (None, None), meaning "header-less mode."
    """
    if anchor_row is not None:
        candidates = range(anchor_row - 1, max(0, anchor_row - 1 - HEADER_SEARCH_WINDOW), -1)
    else:
        candidates = range(1, min(len(sheet.rows), HEADER_SEARCH_WINDOW) + 1)

    for row_index in candidates:
        if row_index < 1 or row_index > len(sheet.rows):
            continue
        row_values = sheet.rows[row_index - 1]
        if _row_looks_like_header(row_values):
            header = [str(v).strip() if v is not None else "" for v in row_values]
            return row_index, header
    return None, None


def row_to_dict(row_values: list[Any], header: list[str] | None) -> dict[str, Any]:
    """Map a row's values onto header names; falls back to positional
    `col_0`, `col_1`, ... keys in header-less mode or for unnamed columns.
    """
    if not header:
        return {f"col_{i}": v for i, v in enumerate(row_values)}
    result: dict[str, Any] = {}
    for i, value in enumerate(row_values):
        key = header[i] if i < len(header) and header[i] else f"col_{i}"
        result[key] = value
    return result
