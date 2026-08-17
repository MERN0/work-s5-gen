from __future__ import annotations

from openpyxl import load_workbook

from app.core.artifacts.swe6.config import Swe6Config
from app.core.artifacts.swe6.models.requirement import EnrichedRequirement, FeatureMeta, Requirement
from app.core.artifacts.swe6.models.test_case import ActionStep, TestCase, VerificationStep
from app.core.artifacts.swe6.output.workbook_builder import build_workbook


def _config(tmp_path) -> Swe6Config:
    out = tmp_path / "out"
    out.mkdir()
    return Swe6Config(output_dir=out, input_folder=tmp_path, req_filename="", req_sheet_name="005")


def _state(tmp_path):
    requirement = Requirement(
        req_id="R1",
        source_file="reqs.xlsx",
        sheet_name="005",
        row_index=10,
        cells={},
        requirement_text="When VehIgn is Active then...",
        matched_keyword="sw qualification test",
        feature=FeatureMeta(feature_group="Ramp / Slope Control", feature_name="005 - Ramp / Slope Control"),
    )
    enriched = EnrichedRequirement(requirement=requirement, supporting_context=[], allowed_signal_tokens=set())

    test_case = TestCase(
        testcase_id="TMHC_SWQTC_001",
        feature="005 - Ramp / Slope Control",
        variant="A1",
        requirement_ids=["R1"],
        test_type="Normal_system",
        objective="Verify slope assist activates",
        description="Test slope assist activation",
        pre_condition="Vehicle powered on",
        input_data=None,
        setup_action_steps=[ActionStep(verb="SET", detail="VehIgn to ON")],
        verification_steps=[
            VerificationStep(verb="VERIFY", detail="SlopeAssistState", expected_result="Active")
        ],
        mode_of_execution="Automated",
        priority="P2",
        odc_trigger=None,
        remarks=None,
    )

    return {
        "config": _config(tmp_path),
        "enriched_requirements": {"R1": enriched},
        "results": [test_case],
        "errors": [],
    }


def test_build_workbook_produces_expected_sheets_and_content(tmp_path):
    state = _state(tmp_path)
    config = state["config"]

    path = build_workbook(state, config)

    assert path.exists()
    wb = load_workbook(path)
    assert set(wb.sheetnames) == {"005", "Cover Page", "Test Pattern", "Item List", "Configurable Parameters"}

    ws = wb["005"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "TestCase ID" in header
    assert "Test Steps" in header
    assert "Expected Results" in header

    data_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    row_dict = dict(zip(header, data_row))
    assert row_dict["TestCase ID"] == "TMHC_SWQTC_001"
    assert row_dict["Test Steps"].startswith("1. Test_start")
    assert "End_of_test" in row_dict["Test Steps"]

    items_ws = wb["Item List"]
    items_header = [c.value for c in next(items_ws.iter_rows(min_row=1, max_row=1))]
    assert items_header == [
        "Sr#",
        "Test type",
        "Feature Name",
        "Test Case ID",
        "Applicable Variants",
        "Execution Required",
        "Remarks",
    ]
    items_row = [c.value for c in next(items_ws.iter_rows(min_row=2, max_row=2))]
    assert items_row[3] == "TMHC_SWQTC_001"


def test_build_workbook_falls_back_gracefully_with_no_matching_supporting_sheets(tmp_path):
    state = _state(tmp_path)
    config = state["config"]

    path = build_workbook(state, config)
    wb = load_workbook(path)

    # No Test Pattern / Configurable Parameters sheet existed in input_folder,
    # so these must be labeled placeholders, not crash the build.
    assert wb["Test Pattern"]["A1"].value == "Test Pattern"
    assert wb["Configurable Parameters"]["A1"].value == "Configurable Parameters"
