"""Synthetic fixture mirroring the real Communication Matrix sheet layout:
Protocol:/Baudrate: metadata sitting well above the real header (row 25 in
the real workbook), with signal rows below it.
"""
from __future__ import annotations

import openpyxl
import pytest

from app.core.artifacts.sys5.config import Sys5Config
from app.core.artifacts.sys5.excel.supporting_docs_loader import load_supporting_entities


@pytest.fixture
def comm_matrix_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BEV Response Network Rel"

    ws.cell(row=22, column=1, value="Protocol:")
    ws.cell(row=22, column=2, value="2.x")
    ws.cell(row=23, column=1, value="Baudrate:")
    ws.cell(row=23, column=2, value="19.2kb/s")

    headers = [
        "TYPE", "Node", "Frame ID", "Frame Name", "Signal Byte No.", "Signal Bit No.",
        "Signal Name", "Signal Function", "Signal length (Bit)", "Init Value",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(row=25, column=col, value=header)

    row_a = ["Fan Response", "FanA", "0x1E", "FanAResp", 0, 0, "FanARespRPM", "Fan Actual Speed", 7, "0x00"]
    row_b = ["Fan Response", "FanA", "0x1E", "FanAResp", 0, 7, "FanARespVolt", "Fan Input voltage", 8, "0x00"]
    for col, value in enumerate(row_a, start=1):
        ws.cell(row=26, column=col, value=value)
    for col, value in enumerate(row_b, start=1):
        ws.cell(row=27, column=col, value=value)

    path = tmp_path / "comm_matrix.xlsx"
    wb.save(path)
    return path


def _config(tmp_path, **overrides) -> Sys5Config:
    defaults = dict(output_dir=tmp_path / "out", input_folder=tmp_path, req_filename="", req_sheet_name="")
    defaults.update(overrides)
    return Sys5Config(**defaults)


def test_load_supporting_entities_skips_metadata_rows_and_finds_real_header(tmp_path, comm_matrix_workbook):
    config = _config(tmp_path)
    entities = load_supporting_entities(config)

    signal_names = {e.fields.get("Signal Name") for e in entities}
    assert "FanARespRPM" in signal_names
    assert "FanARespVolt" in signal_names
    assert not any("Protocol:" in e.searchable_text for e in entities)


def test_load_supporting_entities_guesses_entity_type_from_sheet_name(tmp_path, comm_matrix_workbook):
    config = _config(tmp_path)
    entities = load_supporting_entities(config)
    assert entities
    assert all(e.entity_type == "comm_matrix" for e in entities)


def test_load_supporting_entities_skips_corrupt_files_without_raising(tmp_path, comm_matrix_workbook):
    (tmp_path / "corrupt.xlsx").write_text("not actually a zip/xlsx file")
    config = _config(tmp_path)
    entities = load_supporting_entities(config)
    # the good file's entities still come through; the corrupt one is skipped
    assert any(e.fields.get("Signal Name") == "FanARespRPM" for e in entities)
