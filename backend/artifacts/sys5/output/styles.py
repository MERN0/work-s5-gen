"""Shared openpyxl styling so all 5 sheets look consistent, without every
sheet writer reimplementing header fills/borders/widths.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")
TOP_ALIGNMENT = Alignment(vertical="top")

_thin_side = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=_thin_side, right=_thin_side, top=_thin_side, bottom=_thin_side)


def style_header_row(ws: Worksheet, row: int, num_columns: int) -> None:
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def style_data_row(ws: Worksheet, row: int, num_columns: int, wrap_columns: set[int] | None = None) -> None:
    wrap_columns = wrap_columns or set()
    for col in range(1, num_columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.alignment = WRAP_ALIGNMENT if col in wrap_columns else TOP_ALIGNMENT
        cell.border = THIN_BORDER


def set_column_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
