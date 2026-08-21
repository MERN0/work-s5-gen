"""Cover Page: title/description, a small doc-metadata table, and a Revision
History table with one auto-generated row per run. Always templated from
config + feature metadata — no external template file (per project
decision).
"""
from __future__ import annotations

from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from artifacts.swe6.output.styles import set_column_widths, style_data_row, style_header_row

_REVISION_HEADERS = ["Revision No", "Change Description", "Date", "Created by", "Approved by", "Description"]


def write_cover_page_sheet(ws: Worksheet, cover_data: dict) -> None:
    ws.cell(row=1, column=2, value=cover_data.get("title", ""))
    ws.cell(row=1, column=2).font = Font(bold=True, size=14)

    row = 3
    if cover_data.get("feature_group"):
        ws.cell(row=row, column=2, value=f"Feature Group: {cover_data['feature_group']}")
        row += 1
    if cover_data.get("object_heading"):
        ws.cell(row=row, column=2, value=f"Object Heading: {cover_data['object_heading']}")
        row += 1

    row += 1
    for key, value in cover_data.get("doc_meta", {}).items():
        ws.cell(row=row, column=2, value=key)
        ws.cell(row=row, column=3, value=value)
        row += 1

    row += 2
    ws.cell(row=row, column=2, value="Revision History")
    row += 1
    header_row = row
    for col_index, header in enumerate(_REVISION_HEADERS, start=2):
        ws.cell(row=header_row, column=col_index, value=header)
    style_header_row(ws, row=header_row, num_columns=len(_REVISION_HEADERS) + 1)

    for offset, revision in enumerate(cover_data.get("revision_rows", []), start=header_row + 1):
        for col_index, header in enumerate(_REVISION_HEADERS, start=2):
            ws.cell(row=offset, column=col_index, value=revision.get(header, ""))
        style_data_row(ws, row=offset, num_columns=len(_REVISION_HEADERS) + 1, wrap_columns={3})

    set_column_widths(ws, [4, 24, 30, 16, 16, 16, 30])
