"""Shared writer for sheets that either mirror a matching supporting-doc
sheet found by the format profile, or fall back to a labeled placeholder
when no matching sheet was found in this project's input files.
"""
from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from artifacts.swe6.output.styles import set_column_widths, style_data_row, style_header_row


def write_sourced_table_sheet(ws: Worksheet, data: dict, title: str) -> None:
    if not data.get("found"):
        ws.cell(row=1, column=1, value=title)
        ws.cell(
            row=2,
            column=1,
            value=f"No matching supporting-doc sheet found for {title.lower()} in this project's input files.",
        )
        set_column_widths(ws, [70])
        return

    columns = data["columns"]
    ws.cell(row=1, column=1, value=f"{title} (source: {data.get('source_sheet', '')})")
    header_row = 2
    for col_index, header in enumerate(columns, start=1):
        ws.cell(row=header_row, column=col_index, value=header or f"col_{col_index}")
    style_header_row(ws, row=header_row, num_columns=len(columns))

    for offset, row_dict in enumerate(data["rows"], start=header_row + 1):
        for col_index, header in enumerate(columns, start=1):
            ws.cell(row=offset, column=col_index, value=row_dict.get(header, ""))
        style_data_row(ws, row=offset, num_columns=len(columns))

    set_column_widths(ws, [18] * len(columns))
