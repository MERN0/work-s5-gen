"""Assembles the final 5-sheet .xlsx and writes it to output_dir — the ONLY
file this pipeline may leave there, since runner.py's zip step zips
every non-.zip file it finds sitting in output_dir.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl

from artifacts.swe6.config import Swe6Config
from artifacts.swe6.excel.supporting_docs_loader import load_supporting_sheets
from artifacts.swe6.formats.base import WorkbookBuildContext
from artifacts.swe6.formats.registry import get_format_profile
from artifacts.swe6.models.state import PipelineState
from artifacts.swe6.output.sheets.configurable_parameters_sheet import (
    write_configurable_parameters_sheet,
)
from artifacts.swe6.output.sheets.cover_page_sheet import write_cover_page_sheet
from artifacts.swe6.output.sheets.item_list_sheet import write_item_list_sheet
from artifacts.swe6.output.sheets.test_cases_sheet import write_test_cases_sheet
from artifacts.swe6.output.sheets.test_pattern_sheet import write_test_pattern_sheet

_INVALID_SHEET_CHARS = set('[]:*?/\\')
_INVALID_FILENAME_CHARS = set('<>:"/\\|?*')


def build_workbook(state: PipelineState, config: Swe6Config) -> Path:
    profile = get_format_profile(config.format_profile)
    supporting_sheets = load_supporting_sheets(config)
    ctx = WorkbookBuildContext(state=state, supporting_sheets=supporting_sheets)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    test_cases_sheet_name = config.req_sheet_name or "Test Cases"

    ws_test_cases = wb.create_sheet(_safe_sheet_name(test_cases_sheet_name))
    write_test_cases_sheet(ws_test_cases, state["results"], profile)

    ws_cover = wb.create_sheet("Cover Page")
    write_cover_page_sheet(ws_cover, profile.build_cover_page(ctx))

    ws_pattern = wb.create_sheet("Test Pattern")
    write_test_pattern_sheet(ws_pattern, profile.build_test_pattern(ctx))

    ws_items = wb.create_sheet("Item List")
    write_item_list_sheet(ws_items, state["results"])

    ws_params = wb.create_sheet("Configurable Parameters")
    write_configurable_parameters_sheet(ws_params, profile.build_configurable_parameters(ctx))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = _safe_file_name(f"swe6_{test_cases_sheet_name}_{timestamp}.xlsx")
    output_path = config.output_dir / file_name
    wb.save(output_path)
    return output_path


def _safe_sheet_name(name: str) -> str:
    cleaned = "".join(c for c in name if c not in _INVALID_SHEET_CHARS)
    return cleaned[:31] or "Sheet1"


def _safe_file_name(name: str) -> str:
    return "".join(c for c in name if c not in _INVALID_FILENAME_CHARS)
